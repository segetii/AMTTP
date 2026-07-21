from docx import Document
from docx.oxml.ns import qn
from docx2pdf import convert
import PyPDF2

path = r'C:\amttp\research\visa\SUBMISSION_ORGANISED_BY_CATEGORY\01_Mandatory_Criteria_MC\MC-2_GitHub_Contributions.docx'
doc = Document(path)

# --- 1. Fix Tables 1 & 2 font: force all w:sz to 20 (10pt) ---
# Tables 1 and 2 are doc.tables[1] and doc.tables[2]
for tbl_idx in (1, 2):
    tbl = doc.tables[tbl_idx]
    for el in tbl._tbl.iter(qn('w:sz')):
        el.set(qn('w:val'), '20')  # 10pt
    for el in tbl._tbl.iter(qn('w:szCs')):
        el.set(qn('w:val'), '20')
    print(f'Table {tbl_idx} font fixed to 10pt')

# --- 2. Remove dangling Figure 2 caption (P[23]) and the extra blank before it ---
# Current structure: P[20]=section heading, P[21]=blank, P[22]=blank, P[23]=Figure 2 caption
# Remove P[22] (extra blank) and P[23] (dangling caption) — leaving one blank after heading
body = doc.element.body
paras = doc.paragraphs

p22 = paras[22]
p23 = paras[23]

print(f'Removing P[22]: "{p22.text}"')
print(f'Removing P[23]: "{p23.text[:70]}"')

body.remove(p22._element)
# After removing p22, p23 is now p22
body.remove(doc.paragraphs[22]._element)

print('Removed dangling caption and extra blank')
print('Paragraphs now:', len(doc.paragraphs))

doc.save(path)
print('Saved')

convert(path, path.replace('.docx', '.pdf'))
pdf = path.replace('.docx', '.pdf')
reader = PyPDF2.PdfReader(pdf)
print(f'PDF exported: {len(reader.pages)} pages')
