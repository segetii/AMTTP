from docx import Document
from docx.oxml.ns import qn

path = r'C:\amttp\research\visa\SUBMISSION_ORGANISED_BY_CATEGORY\01_Mandatory_Criteria_MC\MC-2_GitHub_Contributions.docx'
doc = Document(path)
print(f'Paragraphs: {len(doc.paragraphs)}, Tables: {len(doc.tables)}')
print()

body = doc.element.body
para_idx = 0
tbl_idx = 0
for child in body:
    tag = child.tag.split('}')[-1]
    if tag == 'p':
        txt = ''.join(t.text or '' for t in child.iter(qn('w:t')))
        has_img = 'graphicData' in child.xml
        flag = ' [IMG]' if has_img else ''
        sizes = set()
        for rpr in child.findall('.//' + qn('w:rPr')):
            sz = rpr.find(qn('w:sz'))
            if sz is not None:
                sizes.add(int(sz.get(qn('w:val'), 0)) // 2)
        size_str = f' font={sizes}' if sizes else ''
        txt_show = txt[:70] if txt.strip() else '[BLANK]'
        print(f'  P[{para_idx:2d}] {txt_show}{flag}{size_str}')
        para_idx += 1
    elif tag == 'tbl':
        rows = child.findall('.//' + qn('w:tr'))
        print(f'  TABLE {tbl_idx} ({len(rows)} rows)')
        tbl_idx += 1

print()
print('=== IMAGES ===')
for p in doc.paragraphs:
    if 'graphicData' in p._element.xml:
        for ext in p._element.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}ext'):
            cx = int(ext.get('cx', 0))
            cy = int(ext.get('cy', 0))
            print(f'  Image: {cx/914400:.2f}" x {cy/914400:.2f}"')

print()
print('=== DEFAULT STYLE FONT ===')
from docx.shared import Pt
normal = doc.styles['Normal']
if normal.font.size:
    print(f'Normal style font: {normal.font.size.pt}pt')
else:
    print('Normal style font: (inherited/unset)')
