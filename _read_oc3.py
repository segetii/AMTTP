from docx import Document
from docx.oxml.ns import qn

path = r'C:\amttp\research\visa\Evidence Upgrade V2\00_CURRENT_COPY\OC3-1_ML_Pipeline.docx'
doc = Document(path)
print(f'Paragraphs: {len(doc.paragraphs)}, Tables: {len(doc.tables)}')
print()

body = doc.element.body
para_idx = 0
tbl_idx = 0
for child in body:
    tag = child.tag.split('}')[-1]
    if tag == 'p':
        txt = ''.join(t.text or '' for t in child.iter(qn('w:t')))
        has_img = 'graphicData' in child.xml
        flag = ' [IMG]' if has_img else ''
        txt_show = txt if txt.strip() else '[BLANK]'
        print(f'P[{para_idx:2d}] {txt_show[:120]}{flag}')
        para_idx += 1
    elif tag == 'tbl':
        print(f'--- TABLE {tbl_idx} ---')
        for row in child.findall('.//' + qn('w:tr')):
            cells = []
            for cell in row.findall('.//' + qn('w:tc')):
                cell_txt = ''.join(t.text or '' for t in cell.iter(qn('w:t')))
                cells.append(cell_txt[:60])
            print(f'  {cells}')
        tbl_idx += 1
