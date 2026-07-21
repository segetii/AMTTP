"""
Insert a verification table after paragraph [5] (What this proves...)
and before paragraph [6] (Live Public Deployment Evidence).
Then re-export and verify still 3 pages.
"""
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
import copy, fitz, win32com.client, pythoncom

FOLDER = Path(r'C:\amttp\research\visa\SUBMISSION_ORGANISED_BY_CATEGORY\01_Mandatory_Criteria_MC')
SRC = FOLDER / 'MC-1_AMTTP_Live_Demonstration_v2.docx'
PDF = FOLDER / 'MC-1_AMTTP_Live_Demonstration_v2.pdf'

NAVY  = RGBColor(0x1a, 0x2a, 0x4a)
WHITE = RGBColor(0xff, 0xff, 0xff)
LGRAY = RGBColor(0xf0, 0xf4, 0xf8)

doc = Document(str(SRC))

# ── Build the table ───────────────────────────────────────────────────────────
ROWS = [
    # (Claim, Evidence, Independent verification)
    ('Claim', 'Evidence', 'Independent verification'),  # header
    ('Public product exists',  'amttp.com live deployment',
     'Cloudflare analytics + public URL'),
    ('Public interest',        '1,220 unique visitors; 38,320 requests',
     'Cloudflare independent infrastructure analytics'),
    ('Working system',         '4-minute demo and product screens',
     'YouTube demo + War Room screenshots'),
    ('External validation',    'Signed/sealed pilot agreement',
     'Glitterati Estates Schedule A + signature/seal'),
]

tbl = doc.add_table(rows=len(ROWS), cols=3)
tbl.style = 'Table Grid'

# Set column widths: usable = 17.9cm → 4.5 | 7.0 | 6.4
col_widths = [Cm(4.5), Cm(7.0), Cm(6.4)]
for row in tbl.rows:
    for idx, cell in enumerate(row.cells):
        cell.width = col_widths[idx]

for r_idx, (c1, c2, c3) in enumerate(ROWS):
    row = tbl.rows[r_idx]
    is_header = r_idx == 0
    for c_idx, text in enumerate([c1, c2, c3]):
        cell = row.cells[c_idx]
        cell.text = ''  # clear default
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(text)
        run.bold = is_header
        run.font.size = Pt(9)
        if is_header:
            run.font.color.rgb = WHITE
        else:
            run.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)

        # Cell padding
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for side in ('w:top', 'w:bottom', 'w:left', 'w:right'):
            m = OxmlElement(side)
            m.set(qn('w:w'), '60')
            m.set(qn('w:type'), 'dxa')
            tcMar.append(m)
        tcPr.append(tcMar)

        # Header row: navy background
        if is_header:
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'),   'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'),  '1a2a4a')
            tcPr.append(shd)
        elif r_idx % 2 == 0:
            # Alternating light-blue tint
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'),   'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'),  'f0f4f8')
            tcPr.append(shd)

# ── Move the table to after paragraph [5] ────────────────────────────────────
# The new table is currently appended at the end of the body.
# We need to move its XML element to after paragraphs[5]._element.
body = doc.element.body
tbl_elem = tbl._tbl

# Find paragraphs[5] element
target_para_elem = doc.paragraphs[5]._element

# Remove tbl from current position (end) and insert after target
body.remove(tbl_elem)
target_para_elem.addnext(tbl_elem)

TMP = FOLDER / '~tmp2_MC-1_AMTTP_Live_Demonstration_v2.docx'
doc.save(str(TMP))
print('Saved to temp.')

# ── Export + count ────────────────────────────────────────────────────────────
pythoncom.CoInitialize()
word = win32com.client.DispatchEx('Word.Application')
word.Visible = False
word.DisplayAlerts = 0
try:
    d = word.Documents.Open(str(TMP), ReadOnly=True)
    d.SaveAs2(str(PDF), FileFormat=17)
    d.Close(False)
finally:
    word.Quit()

pages = fitz.open(str(PDF)).page_count
print(f'PDF pages: {pages}  {"✓ OK" if pages <= 3 else "OVER"}')

import os
if pages <= 3:
    try:
        os.replace(str(TMP), str(SRC))
        print('Replaced original.')
    except PermissionError:
        print(f'Close Word first, then rename:\n  {TMP.name}\n→ {SRC.name}')
