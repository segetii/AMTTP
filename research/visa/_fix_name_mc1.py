"""
Fix company name in both table and body text, then compile to PDF.
Works from latest temp file (tmp2) which has the table.
"""
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
import fitz, win32com.client, pythoncom, os

FOLDER = Path(r'C:\amttp\research\visa\SUBMISSION_ORGANISED_BY_CATEGORY\01_Mandatory_Criteria_MC')
TMP2 = FOLDER / '~tmp2_MC-1_AMTTP_Live_Demonstration_v2.docx'
OUT  = FOLDER / 'MC-1_AMTTP_Live_Demonstration_v2.docx'
PDF  = FOLDER / 'MC-1_AMTTP_Live_Demonstration_v2.pdf'

doc = Document(str(TMP2))

OLD = 'Glitterati Estates'
NEW = 'Ibironke Estate Limited (RC No. 1647148)'

fixes = 0

# Fix in paragraphs
for para in doc.paragraphs:
    if OLD in para.text:
        for run in para.runs:
            if OLD in run.text:
                run.text = run.text.replace(OLD, NEW)
                fixes += 1

# Fix in tables
for tbl in doc.tables:
    for row in tbl.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    if OLD in run.text:
                        run.text = run.text.replace(OLD, NEW)
                        fixes += 1

print(f'Fixed {fixes} occurrence(s) of company name.')

# Save to a new temp, then replace OUT
TMP3 = FOLDER / '~tmp3_MC-1_AMTTP_Live_Demonstration_v2.docx'
doc.save(str(TMP3))

# Export to PDF
pythoncom.CoInitialize()
word = win32com.client.DispatchEx('Word.Application')
word.Visible = False
word.DisplayAlerts = 0
try:
    d = word.Documents.Open(str(TMP3), ReadOnly=True)
    d.SaveAs2(str(PDF), FileFormat=17)
    d.Close(False)
finally:
    word.Quit()

pages = fitz.open(str(PDF)).page_count
print(f'PDF pages: {pages}  {"✓ OK" if pages <= 3 else "OVER"}')

# Replace the main docx
try:
    os.replace(str(TMP3), str(OUT))
    print(f'Saved: {OUT.name}')
except PermissionError:
    print(f'Close the file in Word first, then rename:\n  {TMP3.name}\n→ {OUT.name}')
