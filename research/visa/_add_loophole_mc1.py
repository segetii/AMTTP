from docx import Document
from docx.shared import Pt, RGBColor

DOC = r'C:\amttp\research\visa\SUBMISSION_ORGANISED_BY_CATEGORY\01_Mandatory_Criteria_MC\MC-1_AMTTP_Live_Demonstration_v2.docx'

doc = Document(DOC)

p = doc.add_paragraph()
r1 = p.add_run('Anti-loophole note: ')
r1.bold = True
r1.font.size = Pt(8)
r1.font.color.rgb = RGBColor(0x1a, 0x2a, 0x4a)

r2 = p.add_run(
    'This exhibit does not rely only on screenshots. It ties the product to third-party '
    'infrastructure analytics, a public URL, a video demonstration, and a signed/sealed '
    'external pilot document.'
)
r2.bold = False
r2.font.size = Pt(8)

doc.save(DOC)
print('Saved OK')
