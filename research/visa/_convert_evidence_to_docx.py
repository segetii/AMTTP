from pathlib import Path
from pypdf import PdfReader
from docx import Document
import shutil

root = Path(r"C:\amttp\research\visa\SUBMISSION_ORGANISED_BY_CATEGORY")
out_root = Path(r"C:\amttp\research\visa\EVIDENCE_DOCX_COPIES")
out_root.mkdir(parents=True, exist_ok=True)

# Keep original category structure
for category in root.iterdir():
    if not category.is_dir():
        continue
    out_cat = out_root / category.name
    out_cat.mkdir(parents=True, exist_ok=True)

    # Copy existing docx directly
    for docx_file in category.glob("*.docx"):
        shutil.copy2(docx_file, out_cat / docx_file.name)

    # Convert each PDF to docx (text-extracted, editable)
    for pdf_file in category.glob("*.pdf"):
        reader = PdfReader(str(pdf_file))
        doc = Document()
        doc.add_heading(pdf_file.stem, level=1)
        doc.add_paragraph(f"Source PDF: {pdf_file.name}")
        doc.add_paragraph("Note: Auto-converted to editable text. Review formatting.")

        for i, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            doc.add_heading(f"Page {i}", level=2)
            # preserve paragraph breaks roughly
            blocks = [b.strip() for b in text.split("\n\n") if b.strip()]
            if not blocks:
                lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
                blocks = ["\n".join(lines)] if lines else ["(No extractable text on this page)"]
            for b in blocks:
                doc.add_paragraph(b)

        out_file = out_cat / (pdf_file.stem + ".docx")
        doc.save(str(out_file))

print("DONE")
print(f"Output folder: {out_root}")
