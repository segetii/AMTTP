from docx import Document
from docx.shared import Pt, RGBColor, Cm

NAVY = RGBColor(0x1a, 0x2a, 0x4a)
GRAY = RGBColor(0x44, 0x44, 0x44)
DOC  = r'C:\amttp\research\visa\SUBMISSION_ORGANISED_BY_CATEGORY\01_Mandatory_Criteria_MC\MC-1_AMTTP_Live_Demonstration.docx'
IMG  = r'C:\amttp\research\visa\images\pilot_agreement_signed.png'

doc = Document(DOC)

# Usable width: 21cm - 1.55cm left - 1.55cm right = 17.9cm
IMG_W = Cm(17.5)

def add_section_heading(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(9.5)
    run.font.color.rgb = NAVY
    return p

def add_body(doc, bold_part, normal_part=''):
    p = doc.add_paragraph()
    if bold_part:
        r = p.add_run(bold_part)
        r.bold = True
        r.font.size = Pt(8)
    if normal_part:
        r = p.add_run(normal_part)
        r.bold = False
        r.font.size = Pt(8)
    return p

def add_caption(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = False
    run.font.size = Pt(7)
    run.font.color.rgb = GRAY
    return p

# Spacer
doc.add_paragraph()

# Section heading
add_section_heading(doc, 'External Validation: Enterprise Pilot Agreement')

# Body paragraphs
add_body(doc, 'What this proves: ',
         'External enterprise adoption under a formal legal agreement, independently verifying that '
         'AMTTP is usable by a real organisation outside the developer\'s own environment. '
         'This is not a personal project or academic prototype — a Nigerian property-technology company '
         'signed a binding pilot evaluation agreement to test AMTTP components in its operational systems.')

add_body(doc, 'Signed parties: ',
         'Olusegun Odeyemi (Developer) and Akintunde Ibironke, Director of Ibironke Estate Limited '
         '(RC No. 1647148). The agreement grants the company a 90-day evaluation licence, '
         'defines IP ownership (retained by Developer), sets confidentiality obligations, and '
         'includes a commercialisation pathway.')

add_body(doc, 'Schedule A confirms: ',
         'The evaluated components — Smart Contract Gateway, ML Risk Engine, Compliance Orchestrator, '
         'War Room Dashboard, and Cross-Chain SDK — are the exact same components evidenced elsewhere '
         'in this submission. The stamped company seal (red) and both signatures are visible below.')

# Image centred
p_img = doc.add_paragraph()
p_img.alignment = 1  # WD_ALIGN_PARAGRAPH.CENTER
run_img = p_img.add_run()
run_img.add_picture(IMG, width=IMG_W)

# Caption
add_caption(doc,
    'Figure 5: Pilot Evaluation Agreement — Schedule A (left) listing the evaluated AMTTP components, '
    'and the signed and company-sealed signature page (right). Signed by Developer and Akintunde Ibironke, '
    'Director, Ibironke Estate Limited. Company seal (red stamp) confirms formal corporate execution.')

OUT = r'C:\amttp\research\visa\SUBMISSION_ORGANISED_BY_CATEGORY\01_Mandatory_Criteria_MC\MC-1_AMTTP_Live_Demonstration_v2.docx'
doc.save(OUT)
print('Saved OK:', OUT)
