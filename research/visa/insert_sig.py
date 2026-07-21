"""
Insert Segun Odeyemi's signature into the Pilot Evaluation Agreement.
Also fixes:
  - Date brackets in P1: [05 March 2026] -> 05 March 2026
  - Developer address in P3: [ADDRESS] -> actual address
  - Restructures P95 to match company sig format: Signature: [IMAGE] Name: ...
  - Inserts Title/Date paragraph after P95
"""
import copy
from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from lxml import etree

SRC = r"C:\Users\Administrator\Documents\Segun Odeyemi's PILOT EVALUATION AGREEMENT - SIGNED.docx"
SIG = r"C:\Users\Administrator\Documents\segun_sig.png"
OUT = r"C:\Users\Administrator\Documents\Segun Odeyemi's PILOT EVALUATION AGREEMENT - BOTH SIGNED.docx"
ADDRESS = "Apartment 714, Prosperity House, Gower Street, Derby, UK, DE1 1AW"
DATE = "14 May 2026"

def make_text_run(text, rpr=None):
    """Create a <w:r> element with the given text."""
    r = OxmlElement('w:r')
    if rpr is not None:
        r.append(copy.deepcopy(rpr))
    t = OxmlElement('w:t')
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    t.text = text
    r.append(t)
    return r

doc = Document(SRC)

# ── Fix P1: remove brackets from date (brackets are in separate runs) ───────
p1 = doc.paragraphs[1]
for run in p1.runs:
    if run.text.endswith('['):
        run.text = run.text[:-1]
    if run.text.startswith('] ') or run.text == ']':
        run.text = run.text.replace('] ', ' ').replace(']', '')
    # Also handle if all in one run
    if '[05 March 2026]' in run.text:
        run.text = run.text.replace('[05 March 2026]', '05 March 2026')
print(f'P1 fixed: {p1.text[:80]}')

# ── Fix P3: fill in developer address ───────────────────────────────────────
p3 = doc.paragraphs[3]
for run in p3.runs:
    if '[ADDRESS]' in run.text:
        run.text = run.text.replace('[ADDRESS]', ADDRESS)
        print(f'P3 fixed: {run.text[:100]}')

# ── Restructure P95: Signature: [IMAGE]   Name: ... ─────────────────────────
p95 = doc.paragraphs[95]
p95_elem = p95._p

# Get the paragraph's run rPr for consistent formatting
pPr = p95_elem.find(qn('w:pPr'))
rpr = pPr.find(qn('w:rPr')) if pPr is not None else None

# Remove all existing <w:r> and <w:ins> runs from p95
for child in list(p95_elem):
    tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
    if tag in ('r', 'ins', 'del', 'bookmarkStart', 'bookmarkEnd'):
        p95_elem.remove(child)

# Add picture to a temp paragraph at end of doc, grab the drawing element
temp_para = doc.add_paragraph()
temp_run = temp_para.add_run()
temp_run.add_picture(SIG, width=Inches(1.5))
drawing_elem = temp_run._r.find(qn('w:drawing'))

# Build the three runs for P95
r1 = make_text_run('Signature: ', rpr)
p95_elem.append(r1)

r_img = OxmlElement('w:r')
r_img.append(copy.deepcopy(drawing_elem))
p95_elem.append(r_img)

r3 = make_text_run('     Name: Mr Olusegun Odeyemi', rpr)
p95_elem.append(r3)

# Remove temp paragraph
temp_para._p.getparent().remove(temp_para._p)
print(f'P95 restructured with signature image')

# ── Insert new P96 (Title/Date) after P95, before existing P96 "For the Company" ──
# We need to insert a new paragraph between P95 and the current P96
new_para_elem = OxmlElement('w:p')

# Copy pPr from P95
if pPr is not None:
    new_para_elem.append(copy.deepcopy(pPr))

title_run = make_text_run(f'Title: Developer     Date: {DATE}', rpr)
new_para_elem.append(title_run)

# Insert after p95_elem
p95_elem.addnext(new_para_elem)
print(f'Inserted Title/Date paragraph after P95')

# ── Save ────────────────────────────────────────────────────────────────────
doc.save(OUT)
print(f'\nSaved: {OUT}')

# Verify
doc2 = Document(OUT)
for i in range(93, 103):
    if i < len(doc2.paragraphs):
        t = doc2.paragraphs[i].text
        has_img = any(
            run._r.find('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}inline') is not None
            for run in doc2.paragraphs[i].runs
        )
        print(f'P{i}: {repr(t[:90])}{"  [HAS IMAGE]" if has_img else ""}')
